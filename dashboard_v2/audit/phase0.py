#!/usr/bin/env python3
"""Phase 0 of dashboard numbers audit:
1. Citation index from Indian_REITs_Key_Financials_FILLED.xlsx cell comments.
2. First-pass diff: dashboard reit-data.json fin values vs workbook cell values.
3. Arithmetic cross-checks (no PDFs).
4. Audit-target dumps for workflow agents.
"""
import json, os, re, sys
import openpyxl

BASE = "/Users/ishan/Downloads/INDIAN REITS"
DASH = os.path.join(BASE, "dashboard_v2")
OUT = os.path.dirname(os.path.abspath(__file__))

SHEETS = {
    "embassy": "Embassy", "mindspace": "Mindspace", "brookfield": "Brookfield",
    "nexus": "Nexus", "krt": "Knowledge (KRT)", "bagmane": "Bagmane",
}
# dashboard fin key -> workbook main-table row label
METRIC_MAP = {
    "revenue": "Revenue from Operations",
    "ndcf": "Net Distributable Cash Flow (NDCF)",
    "dist_total": "Total Distribution",
    "dpu": "Distribution per Unit (DPU)",
    "gav": "Gross Asset Value (GAV) / AUM",
    "gross_debt": "Gross Debt",
    "cash": "Cash & Cash Equivalents",
    "networth": "Net Worth / Unitholders Equity",
    "nav": "NAV per Unit",
    "units_mn": "Units Outstanding",
    "price_eoy": "Unit Price (period-end)",
    "ltv": "Loan-to-Value (LTV)",
    "cost_debt": "Cost of Debt",
    "msf_total": "Total Leasable Area (GLA)",
    "msf_op": "Operational Area",
}

wb = openpyxl.load_workbook(os.path.join(BASE, "Indian_REITs_Key_Financials_FILLED.xlsx"), data_only=True)
reit_data = json.load(open(os.path.join(DASH, "public/data/reit-data.json")))
val_hy = json.load(open(os.path.join(DASH, "public/data/val-hy.json")))
bench = json.load(open(os.path.join(DASH, "public/data/bench.json")))
invit = json.load(open(os.path.join(DASH, "public/data/invit.json")))

# ---------- 1. citation index ----------
citations = []   # rows: reit, table, metric, unit, period, wb_value, source_comment
for rk, sheet in SHEETS.items():
    ws = wb[sheet]
    # locate header rows: any row whose col-A value == 'Metric'
    header_rows = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Metric"]
    for hr in header_rows:
        # each sheet may hold side-by-side CONSOLIDATED / STANDALONE blocks;
        # the band label sits in the row above the header (e.g. row 6), leftmost col of block
        headers = {}
        band = "CONSOLIDATED"
        for c in range(3, ws.max_column + 1):
            for br in (hr - 1, hr - 2):
                bv = ws.cell(br, c).value if br >= 1 else None
                if isinstance(bv, str) and bv.strip() in ("CONSOLIDATED", "STANDALONE"):
                    band = bv.strip()
            hv = ws.cell(hr, c).value
            if hv and str(hv).strip() != "YoY%":
                headers[c] = (str(hv).strip(), band)
        # data rows run until next blank-streak or next header
        r = hr + 1
        blanks = 0
        while r <= ws.max_row and blanks < 4:
            label = ws.cell(r, 1).value
            if label is None:
                blanks += 1; r += 1; continue
            blanks = 0
            if label == "Metric":
                break
            unit = ws.cell(r, 2).value
            for c, (period, band) in headers.items():
                cell = ws.cell(r, c)
                if cell.value is None and cell.comment is None:
                    continue
                citations.append({
                    "reit": rk, "table": "main" if hr < 50 else "detail",
                    "basis": band,
                    "metric": str(label).strip(), "unit": unit,
                    "period": period,
                    "wb_value": cell.value,
                    "source": cell.comment.text.strip() if cell.comment else None,
                })
            r += 1
json.dump(citations, open(os.path.join(OUT, "citations.json"), "w"), indent=1)
print(f"citations.json: {len(citations)} rows ({sum(1 for c in citations if c['source'])} with source comments)")

# ---------- 2. dashboard vs workbook diff ----------
def tol_ok(a, b, metric):
    if a is None or b is None: return None
    try: a, b = float(a), float(b)
    except (TypeError, ValueError): return None
    if metric in ("ltv", "cost_debt"):          # stored as % or fraction — normalise
        if a <= 1 and b > 1: a *= 100
        if b <= 1 and a > 1: b *= 100
        return abs(a - b) <= 0.1
    if metric in ("dpu", "nav", "price_eoy"):   # per-unit ₹
        return abs(a - b) <= 0.05
    if metric in ("msf_total", "msf_op"):
        return abs(a - b) <= 0.1
    if metric == "units_mn":
        return abs(a - b) <= 0.5
    return abs(a - b) <= max(1.0, 0.005 * abs(b))   # ₹cr

cit_lookup = {}
for c in citations:
    if c["table"] == "main" and c["basis"] == "CONSOLIDATED":
        cit_lookup[(c["reit"], c["metric"], c["period"])] = c

diff_rows = []
for rk, fin in reit_data["fin"].items():
    years = fin["years"]
    for mkey, wlabel in METRIC_MAP.items():
        vals = fin.get(mkey)
        if not vals: continue
        for i, fy in enumerate(years):
            dv = vals[i] if i < len(vals) else None
            cit = cit_lookup.get((rk, wlabel, fy))
            wv = cit["wb_value"] if cit else None
            if dv is None and wv is None: continue
            ok = tol_ok(dv, wv, mkey)
            status = ("wb_missing" if cit is None else
                      "dash_null" if dv is None else
                      "wb_null" if wv is None else
                      "match" if ok else "DIFF")
            if status in ("DIFF", "wb_null", "dash_null", "wb_missing"):
                diff_rows.append({"reit": rk, "metric": mkey, "fy": fy,
                                  "dash": dv, "workbook": wv, "status": status,
                                  "source": (cit or {}).get("source")})
json.dump(diff_rows, open(os.path.join(OUT, "dash_vs_workbook.json"), "w"), indent=1)
print(f"dash_vs_workbook.json: {len(diff_rows)} non-matching rows "
      f"({sum(1 for d in diff_rows if d['status']=='DIFF')} DIFF, "
      f"{sum(1 for d in diff_rows if d['status']=='wb_missing')} wb_missing, "
      f"{sum(1 for d in diff_rows if d['status']=='wb_null')} wb_null, "
      f"{sum(1 for d in diff_rows if d['status']=='dash_null')} dash_null)")
for d in diff_rows:
    if d["status"] == "DIFF":
        print(f"  DIFF {d['reit']:10s} {d['metric']:10s} {d['fy']}: dash={d['dash']} wb={d['workbook']}")

# ---------- 3. arithmetic cross-checks ----------
checks = []
def add(reit, name, fy, expect, actual, tol, note=""):
    if expect is None or actual is None: return
    ok = abs(expect - actual) <= tol
    checks.append({"reit": reit, "check": name, "fy": fy,
                   "expected": round(expect, 3), "actual": actual,
                   "ok": ok, "note": note})

for rk, fin in reit_data["fin"].items():
    years = fin["years"]
    for i, fy in enumerate(years):
        g = lambda k: (fin.get(k) or [None]*len(years))[i] if i < len(fin.get(k) or []) else None
        dpu, units, dist = g("dpu"), g("units_mn"), g("dist_total")
        if dpu and units and dist:
            add(rk, "dpu*units≈dist_total", fy, dpu * units / 10, dist, max(2.0, 0.01 * dist),
                "dist ₹cr = dpu ₹ × units mn ÷ 10; small gaps = intra-year unit count changes")
        gd, cash, gav, ltv = g("gross_debt"), g("cash"), g("gav"), g("ltv")
        if gd and gav and ltv:
            nd_ltv = (gd - (cash or 0)) / gav * 100
            gd_ltv = gd / gav * 100
            lv = ltv * 100 if ltv <= 1 else ltv
            best = min(abs(nd_ltv - lv), abs(gd_ltv - lv))
            checks.append({"reit": rk, "check": "ltv_vs_debt/gav", "fy": fy,
                           "expected": f"net {nd_ltv:.1f}% / gross {gd_ltv:.1f}%",
                           "actual": lv, "ok": best <= 1.5,
                           "note": "ltv should be near net- or gross-debt/GAV"})
        nav, networth = g("nav"), g("networth")
        if nav and units and gav and gd:
            approx = (gav - gd + (cash or 0)) / units * 10
            checks.append({"reit": rk, "check": "nav_sanity(gav-debt+cash)/units", "fy": fy,
                           "expected": round(approx, 1), "actual": nav,
                           "ok": abs(approx - nav) <= 0.18 * nav,
                           "note": "loose: ignores other assets/liabilities"})

# spv sums vs latest gav
for rk, assets in reit_data["spv"].items():
    rows = assets if isinstance(assets, list) else assets.get("rows", [])
    tot = sum(a.get("value_cr") or 0 for a in rows if isinstance(a, dict))
    fin = reit_data["fin"].get(rk, {})
    gavs = [v for v in (fin.get("gav") or []) if v]
    latest_gav = gavs[-1] if gavs else None
    if tot and latest_gav:
        checks.append({"reit": rk, "check": "Σspv.value_cr≈latest gav", "fy": fin["years"][len(gavs)-1] if gavs else "?",
                       "expected": latest_gav, "actual": round(tot, 1),
                       "ok": abs(tot - latest_gav) <= 0.05 * latest_gav,
                       "note": "SpvTable warns >5%; spv as-of may differ from FY-end"})

json.dump(checks, open(os.path.join(OUT, "crosschecks.json"), "w"), indent=1)
bad = [c for c in checks if not c["ok"]]
print(f"crosschecks.json: {len(checks)} checks, {len(bad)} failing")
for c in bad:
    print(f"  FAIL {c['reit']:10s} {c['check']:32s} {c['fy']}: expected≈{c['expected']} actual={c['actual']}")

# ---------- 4. audit-target dumps ----------
targets = {"fin": reit_data["fin"], "bv": reit_data.get("bv"),
           "issuances": reit_data.get("issuances"), "blocks": reit_data.get("blocks"),
           "spv": reit_data["spv"], "notes": reit_data.get("notes"),
           "val_hy": val_hy,
           "bench_overview": bench.get("overview"), "bench_metrics": bench.get("metrics"),
           "fd_steps": bench.get("fd_steps"),
           "invit_unverified": invit}
json.dump(targets, open(os.path.join(OUT, "targets.json"), "w"), indent=1)
print("targets.json written")

# quarterly coverage summary
for rk, fin in reit_data["fin"].items():
    q = fin.get("q") or []
    if q:
        qq = [x.get("q") or x.get("label") for x in q] if isinstance(q, list) else list(q.keys())
        print(f"  q[] {rk}: {len(q)} quarters ({qq[0]}..{qq[-1]})")
